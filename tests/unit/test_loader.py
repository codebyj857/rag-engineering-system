from __future__ import annotations

import csv
import json

import pytest
from docx import Document

from rag_engine.ingestion.loader import (
    DocumentExtractionError,
    DocumentLoaderError,
    LoadedDocument,
    UnsupportedDocumentTypeError,
    load_document,
)


def test_txt_loading(tmp_path) -> None:
    path = tmp_path / "sample.txt"
    path.write_text(
        "Hello   world.\n\nThis is a test document.",
        encoding="utf-8",
    )

    document = load_document(path)

    assert isinstance(document, LoadedDocument)
    assert document.source == str(path)
    assert document.filename == "sample.txt"
    assert document.file_type == ".txt"
    assert document.content == "Hello world.\nThis is a test document."
    assert document.metadata == {}


def test_txt_utf8_bom_loading(tmp_path) -> None:
    path = tmp_path / "bom.txt"
    path.write_text(
        "BOM encoded text.",
        encoding="utf-8-sig",
    )

    document = load_document(path)

    assert document.content == "BOM encoded text."


def test_csv_loading(tmp_path) -> None:
    path = tmp_path / "sample.csv"

    with path.open("w", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)
        writer.writerow(["name", "topic"])
        writer.writerow(["Alice", "RAG"])
        writer.writerow(["Bob", "Embeddings"])

    document = load_document(path)

    assert document.file_type == ".csv"
    assert "[Row 1]" in document.content
    assert "name: Alice" in document.content
    assert "topic: RAG" in document.content
    assert "[Row 2]" in document.content
    assert "name: Bob" in document.content
    assert document.metadata["columns"] == ["name", "topic"]
    assert document.metadata["row_count"] == 2


def test_json_loading(tmp_path) -> None:
    path = tmp_path / "sample.json"
    path.write_text(
        json.dumps(
            {
                "topic": "RAG",
                "score": 0.95,
            }
        ),
        encoding="utf-8",
    )

    document = load_document(path)

    assert document.file_type == ".json"
    assert '"topic": "RAG"' in document.content
    assert '"score": 0.95' in document.content
    assert document.metadata["json_type"] == "dict"


def test_docx_loading(tmp_path) -> None:
    path = tmp_path / "sample.docx"

    doc = Document()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Topic"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "RAG"

    doc.save(path)

    document = load_document(path)

    assert document.file_type == ".docx"
    assert "First paragraph." in document.content
    assert "Second paragraph." in document.content
    assert "[Table 1]" in document.content
    assert "Name | Topic" in document.content
    assert "Alice | RAG" in document.content
    assert document.metadata["paragraph_count"] == 2
    assert document.metadata["table_count"] == 1


def test_pdf_loading(tmp_path) -> None:
    import pymupdf

    path = tmp_path / "sample.pdf"

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "PDF test content.")
    pdf.save(path)
    pdf.close()

    document = load_document(path)

    assert document.file_type == ".pdf"
    assert "[Page 1]" in document.content
    assert "PDF test content." in document.content
    assert document.metadata["page_count"] == 1


def test_pdf_metadata_is_preserved(tmp_path) -> None:
    import pymupdf

    path = tmp_path / "metadata.pdf"

    pdf = pymupdf.open()
    pdf.set_metadata(
        {
            "title": "Test PDF",
            "author": "RAG Test",
        }
    )
    page = pdf.new_page()
    page.insert_text((72, 72), "Metadata test.")
    pdf.save(path)
    pdf.close()

    document = load_document(path)

    assert document.metadata["page_count"] == 1
    assert document.metadata["pdf_metadata"]["title"] == "Test PDF"
    assert document.metadata["pdf_metadata"]["author"] == "RAG Test"


def test_unsupported_extension(tmp_path) -> None:
    path = tmp_path / "sample.xyz"
    path.write_text("unsupported", encoding="utf-8")

    with pytest.raises(UnsupportedDocumentTypeError):
        load_document(path)


def test_missing_file() -> None:
    with pytest.raises(FileNotFoundError):
        load_document("does-not-exist.txt")


def test_directory_path_is_rejected(tmp_path) -> None:
    directory = tmp_path / "documents"
    directory.mkdir()

    with pytest.raises(DocumentLoaderError):
        load_document(directory)


def test_empty_text_document_is_rejected(tmp_path) -> None:
    path = tmp_path / "empty.txt"
    path.write_text("   \n\n   ", encoding="utf-8")

    with pytest.raises(DocumentExtractionError):
        load_document(path)